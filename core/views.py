from bson.errors import InvalidId
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from core_requests.models import *
from core_requests.serializers import *
from django.core.mail import send_mail
from django.contrib.auth.hashers import make_password, check_password
from drf_yasg.utils import swagger_auto_schema
from drf_yasg import openapi
from rest_framework.decorators import api_view, permission_classes
from rest_framework.generics import GenericAPIView, ListAPIView
from django.shortcuts import get_object_or_404
from mongoengine.errors import DoesNotExist, ValidationError as MongoValidationError
from rest_framework.exceptions import NotFound
from openpyxl.utils import get_column_letter
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.exceptions import TokenError
from rest_framework_simplejwt.authentication import JWTAuthentication
from rest_framework.permissions import IsAuthenticated
from bson import ObjectId, errors as bson_errors
from bson import ObjectId
from authorization.models import User
from authorization.serializers import UserSerializer
from core_requests.utils import send_push_notification_to_user
from rest_framework.pagination import LimitOffsetPagination
from django.shortcuts import render
from authorization.views import *
from core.serializers import *
from mongoengine import Document, StringField, BooleanField
from collections import Counter, defaultdict
from datetime import datetime
import csv
from collections import defaultdict
from django.http import HttpResponse
import io
import pandas as pd
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.http import FileResponse
from io import BytesIO
from datetime import datetime,timedelta
from drf_yasg.utils import swagger_auto_schema
from drf_yasg import openapi
from core_requests.models import Request  
from core_requests.models import Car, Car_user, Route, Trip
from django.http import JsonResponse
from authorization.models import User
from bson import ObjectId

def test_location_view(request):
    return render(request, 'test-location.html')


class ReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        start_date = request.query_params.get("start")
        end_date = request.query_params.get("end")
        if not start_date or not end_date:
            return Response({"detail": "Start and end dates required."}, status=400)

        try:
            start = datetime.fromisoformat(start_date)
            end = datetime.fromisoformat(end_date)
        except ValueError:
            return Response({"detail": "Invalid date format. Use ISO format."}, status=400)
        requests_in_range = Request.objects(date__gte=start, date__lte=end)
        request_ids = [req.id for req in requests_in_range]
        routes = Route.objects(request__in=request_ids)

        trips = Trip.objects(end_time__gte=start, end_time__lte=end)
        driver_trip_counts = defaultdict(int)
        for trip in trips:
            if trip.car_user and trip.car_user.user:
                driver_trip_counts[str(trip.car_user.user.id)] += 1

        return Response({
            "routes_count": routes.count(),
            "driver_load": driver_trip_counts,
        })


class ReportCSVDownloadView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        role = request.user.role.value
        start_date = request.query_params.get("start")
        end_date = request.query_params.get("end")
        if not start_date or not end_date:
            return Response({"detail": "Start and end dates required."}, status=400)

        try:
            start = datetime.fromisoformat(start_date)
            end = datetime.fromisoformat(end_date)
        except ValueError:
            return Response({"detail": "Invalid date format. Use ISO format."}, status=400)

        requests_in_range = Request.objects(date__gte=start, date__lte=end)
        request_ids = [req.id for req in requests_in_range]
        routes = Route.objects(request__in=request_ids)

        trips = Trip.objects(end_time__gte=start, end_time__lte=end)
        driver_trip_counts = defaultdict(int)
        for trip in trips:
            if trip.car_user and trip.car_user.user:
                driver_trip_counts[str(trip.car_user.user.id)] += 1

        response = HttpResponse(content_type='text/csv')
        filename = f"report_{start.date()}_to_{end.date()}.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        writer = csv.writer(response)
        writer.writerow(['Отчёт по маршрутам и загрузке водителей'])
        writer.writerow(['Период:', f'{start}', '-', f'{end}'])
        writer.writerow([])

        writer.writerow(['Количество маршрутов за период', routes.count()])
        writer.writerow([])

        writer.writerow(['ID водителя', 'Количество поездок'])
        for driver_id, trip_count in driver_trip_counts.items():
            writer.writerow([driver_id, trip_count])

        return response


class CarListCreateAPIView(APIView):
    pagination_class = UnlimitedPagination
    permission_classes = [IsAuthenticated]

    @swagger_auto_schema(
        operation_description="Получить список автомобилей с пагинацией.",
        manual_parameters=[
            openapi.Parameter(
                'limit',
                openapi.IN_QUERY,
                description="Количество машин на странице (по умолчанию 10, максимум не ограничен)",
                type=openapi.TYPE_INTEGER
            ),
            openapi.Parameter(
                'offset',
                openapi.IN_QUERY,
                description="Смещение от начала списка (по умолчанию 0)",
                type=openapi.TYPE_INTEGER
            ),
        ],
        responses={
            200: openapi.Response(
                description="Пагинированный список автомобилей",
                schema=CarSerializer(many=True)
            ),
            400: "Неверные параметры запроса (например, отрицательный limit/offset)"
        }
    )
    def get(self, request):
        cars = Car.objects.all()
        paginator = self.pagination_class()
        result_page = paginator.paginate_queryset(cars, request)
        serializer = CarSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)

    @swagger_auto_schema(
        operation_description="Создать новый автомобиль.",
        request_body=CarSerializer,
        responses={
            201: CarSerializer,
            400: "Ошибка валидации данных"
        }
    )
    def post(self, request):
        serializer = CarSerializer(data=request.data)
        if serializer.is_valid():
            car = serializer.save()
            return Response(CarSerializer(car).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class CarDetailAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @swagger_auto_schema(
        responses={200: CarSerializer()}
    )
    def get_object(self, pk):
        try:
            return Car.objects.get(id=pk)
        except Car.DoesNotExist:
            raise NotFound(detail="Car not found", code=404)

    @swagger_auto_schema(
        request_body=CarSerializer,
        responses={200: CarSerializer()}
    )
    def put(self, request, pk):
        car = self.get_object(pk)
        serializer = CarSerializer(car, data=request.data)
        if serializer.is_valid():
            updated_car = serializer.save()
            return Response(CarSerializer(updated_car).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @swagger_auto_schema(
        responses={204: 'No Content'}
    )
    def delete(self, request, pk):
        car = get_object_or_404(Car, pk=pk)
        car.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class CarUserListCreateAPIView(APIView):
    permission_classes = [IsAuthenticated]
    pagination_class = UnlimitedPagination

    @swagger_auto_schema(
        operation_description="Получить список связей 'Автомобиль-Пользователь' с пагинацией.",
        manual_parameters=[
            openapi.Parameter(
                'limit',
                openapi.IN_QUERY,
                description="Количество элементов на странице (по умолчанию 10, максимум не ограничен)",
                type=openapi.TYPE_INTEGER
            ),
            openapi.Parameter(
                'offset',
                openapi.IN_QUERY,
                description="Смещение от начала списка (по умолчанию 0)",
                type=openapi.TYPE_INTEGER
            ),
        ],
        responses={
            200: openapi.Response(
                description="Пагинированный список связей 'Автомобиль-Пользователь'",
                schema=CarUserSerializer(many=True)
            ),
            400: "Неверные параметры запроса (например, отрицательный limit/offset)"
        }
    )
    def get(self, request):
        car_users = Car_user.objects.all()
        paginator = self.pagination_class()
        result_page = paginator.paginate_queryset(car_users, request)
        serializer = CarUserSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)

    @swagger_auto_schema(
        operation_description="Создать новую связь 'Автомобиль-Пользователь'.",
        request_body=CarUserSerializer,
        responses={
            201: CarUserSerializer,
            400: "Ошибка валидации данных"
        }
    )
    def post(self, request):
        serializer = CarUserSerializer(data=request.data)
        if serializer.is_valid():
            car_user = serializer.save()
            return Response(CarUserSerializer(car_user).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class CarUserDetailAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get_object(self, pk):
        try:
            return Car_user.objects.get(id=pk)
        except Car_user.DoesNotExist:
            raise NotFound("Car_user not found")

    @swagger_auto_schema(responses={200: CarUserSerializer()})
    def get(self, request, pk):
        car_user = self.get_object(pk)
        return Response(CarUserSerializer(car_user).data)

    @swagger_auto_schema(
        request_body=CarUserSerializer,
        responses={200: CarUserSerializer()}
    )
    def put(self, request, pk):
        try:
            car_user = Car_user.objects.get(id=ObjectId(pk))
        except Car_user.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)

        serializer = CarUserSerializer(car_user, data=request.data)
        if serializer.is_valid():
            updated = serializer.save()
            return Response(CarUserSerializer(updated).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @swagger_auto_schema(request_body=CarUserSerializer, responses={200: CarUserSerializer()})
    def patch(self, request, pk):
        try:
            car_user = Car_user.objects.get(id=ObjectId(pk))
        except Car_user.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)

        serializer = CarUserSerializer(car_user, data=request.data, partial=True)
        if serializer.is_valid():
            updated = serializer.save()
            return Response(CarUserSerializer(updated).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @swagger_auto_schema(responses={204: 'No content'})
    def delete(self, request, pk):
        car_user = self.get_object(pk)
        car_user.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


STATUS_TRANSLATION = {
    0: "В процессе",
    1: "Подтверждено",
    2: "Отклонено"
}

class ExportExcelAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @swagger_auto_schema(
        operation_description="Экспорт всех заявок в формате Excel с подробными данными и проверкой машины водителя",
        responses={200: "Файл Excel"}
    )
    def get(self, request):
        data = []

        car_users = {str(cu.user.id): cu for cu in Car_user.objects}

        for req in Request.objects:
            for route in req.routes:
                driver = req.driver
                cu = car_users.get(str(driver.id)) if driver else None

                driver_name = driver.name if driver and cu else "—"
                car_info = cu.car.name if cu and hasattr(cu.car, "name") else "Нет машины"

                row = {
                    "ID пользователя": str(req.user.id) if req.user else "—",
                    "Имя пользователя": req.user.name if req.user else "—",
                    "Имя водителя": driver_name,
                    "Статус": STATUS_TRANSLATION.get(req.status, "Неизвестно"),
                    "Комментарий": req.comments or "",
                    "Дата заявки": req.date.strftime("%Y-%m-%d") if req.date else "",
                    "Откуда": route.departure or "",
                    "Куда": route.destination or "",
                    "Цель поездки": route.goal or "",
                    "Время (длительность)": route.time or "",
                    "Время начала": route.start_time.strftime("%Y-%m-%d %H:%M") if route.start_time else "",
                    "Время окончания": route.end_time.strftime("%Y-%m-%d %H:%M") if route.end_time else "",
                    "Дата поездки": route.travel_date.strftime("%Y-%m-%d") if route.travel_date else "",
                    "Тип транспорта": {
                        "passenger": "пассажирский",
                        "cargo": "грузовой",
                        "light": "легковой"
                    }.get(route.transport_type, ""),
                    "Название машины": car_info,
                }
                data.append(row)

        df = pd.DataFrame(data)

        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Заявки')

            worksheet = writer.sheets['Заявки']
            for i, column in enumerate(df.columns, 1):
                max_length = max(
                    df[column].astype(str).map(len).max(),
                    len(str(column))
                ) + 2  
                worksheet.column_dimensions[get_column_letter(i)].width = max_length

        output.seek(0)
        filename = f"Заявки_{datetime.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"
        response = FileResponse(output, as_attachment=True, filename=filename)
        return response

class DownloadUserActivityReport(APIView):
    @swagger_auto_schema(
        operation_summary="Скачать отчёт по активности пользователей (Excel)",
        manual_parameters=[
            openapi.Parameter(
                'start_date', openapi.IN_QUERY, description="Начальная дата (YYYY-MM-DD)",
                type=openapi.TYPE_STRING, required=False
            ),
            openapi.Parameter(
                'end_date', openapi.IN_QUERY, description="Конечная дата (YYYY-MM-DD)",
                type=openapi.TYPE_STRING, required=False
            )
        ],
        responses={200: 'Файл Excel с отчётом'}
    )
    def get(self, request):
        start_date = request.GET.get('start_date')
        end_date = request.GET.get('end_date')

        users = User.objects(role='user')
        data = []

        for user in users:
            requests = Request.objects(user=user)

            if start_date:
                requests = requests.filter(date__gte=start_date)
            if end_date:
                requests = requests.filter(date__lte=end_date)

            confirmed = requests.filter(status=1).count()
            rejected_requests = requests.filter(status=2)
            rejected = rejected_requests.count()

            # Собираем комментарии из отклонённых заявок
            rejected_comments = []
            for req in rejected_requests:
                if req.comments:
                    rejected_comments.append(req.comments.strip())

            total_time = timedelta()
            for req in requests:
                for route in req.routes:
                    if route.start_time and route.end_time:
                        duration = route.end_time - route.start_time
                        if duration.total_seconds() > 0:
                            total_time += duration

            data.append({
                "ID пользователя": user.id ,
                "Имя пользователя": user.name or user.email,
                "Подразделение": user.subdepartment or "",
                "Создано заявок": requests.count(),
                "Одобрено заявок": confirmed,
                "Отклонено заявок": rejected,
                "Комментарии к отклонённым заявкам": "; ".join(rejected_comments),
                "Общее время поездок": str(total_time),
            })


        df = pd.DataFrame(data)

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
            df.to_excel(writer, index=False, sheet_name='Отчёт')
            worksheet = writer.sheets['Отчёт']
            for i, column in enumerate(df.columns):
                column_width = max(df[column].astype(str).map(len).max(), len(column)) + 2
                worksheet.set_column(i, i, column_width)

        output.seek(0)
        response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="user_activity_report.xlsx"'
        return response
    

class DriverLoadReportExcel(APIView):
    permission_classes = [IsAuthenticated]

    @swagger_auto_schema(
        operation_description="Экспорт отчёта о загруженности водителей",
        responses={200: "Excel-файл"}
    )
    def get(self, request):
        report = defaultdict(lambda: {
            "Кол-во поездок": 0,
            "Кол-во рабочих часов": timedelta(),
            "Общий пробег": 0.0,
            "Пробег без пассажира": 0.0,
            "Кол-во отклонённых заявок": 0
        })

        requests = Request.objects.select_related()
        for req in requests:
            if not req.driver:
                continue
            driver_name = req.driver.name

            if req.status == 1: 
                report[driver_name]["Кол-во поездок"] += 1
                for route in req.routes:
                    if route.start_time and route.end_time:
                        report[driver_name]["Кол-во рабочих часов"] += (route.end_time - route.start_time)
            elif req.status == 2:
                report[driver_name]["Кол-во отклонённых заявок"] += 1

        odometers = OdometerReading.objects.select_related()
        for od in odometers:
            if not od.driver:
                continue
            driver_name = od.driver.name
            if od.start_odometer is not None and od.end_odometer is not None:
                report[driver_name]["Общий пробег"] += abs(od.end_odometer - od.start_odometer)
            if od.no_goal_mileage is not None:
                report[driver_name]["Пробег без пассажира"] += od.no_goal_mileage

        rows = []
        for driver, stats in report.items():
            trips = stats["Кол-во поездок"]
            total_km = stats["Общий пробег"]
            avg_distance = total_km / trips if trips else 0
            working_hours = round(stats["Кол-во рабочих часов"].total_seconds() / 3600, 2)

            rows.append({
                "Имя водителя": driver,
                **stats,
                "Среднее расстояние на поездку": round(avg_distance, 2),
                "Кол-во рабочих часов": working_hours,
            })

        df = pd.DataFrame(rows)

        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name="Загруженность")
            sheet = writer.sheets["Загруженность"]
            for i, column in enumerate(df.columns, 1):
                max_len = max(df[column].astype(str).map(len).max(), len(str(column))) + 2
                sheet.column_dimensions[get_column_letter(i)].width = max_len

        output.seek(0)
        filename = f"Загруженность_водителей_{datetime.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"
        return FileResponse(output, as_attachment=True, filename=filename)


# def get_pair(request):
#     user_id = request.GET.get("user_id")
#     role = request.GET.get("role")
#     if not user_id or not role:
#         return JsonResponse({"error": "user_id and role are required"}, status=400)

#     try:
#         if role == "user":
#             req = Request.objects.get(user=ObjectId(user_id), status=0)
#             return JsonResponse({"user_id": str(req.user.id), "driver_id": str(req.driver.id)})
#         elif role == "driver":
#             req = Request.objects.get(driver=ObjectId(user_id), status=0)
#             return JsonResponse({"user_id": str(req.user.id), "driver_id": str(req.driver.id)})
#         else:
#             return JsonResponse({"error": "Invalid role"}, status=400)
#     except Request.DoesNotExist:
#         return JsonResponse({"error": "No active request found"}, status=404)

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from bson import ObjectId
from django.core.exceptions import ValidationError

class GetPairAPIView(APIView):
    def get(self, request):
        try:
            user_id = request.GET.get("user_id")
            role = request.GET.get("role")

            if not user_id or not role:
                return Response(
                    {"error": "user_id and role are required"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # Проверяем валидность ObjectId
            try:
                user_oid = ObjectId(user_id)
            except:
                return Response(
                    {"error": "Invalid user_id format"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            if role == "user":
                req = Request.objects.get(user=user_oid, status=1)
                response_data = {
                    "user_id": str(req.user.id),
                    "driver_id": str(req.driver.id)
                }
            elif role == "driver":
                req = Request.objects.get(driver=user_oid, status=1)
                response_data = {
                    "user_id": str(req.user.id),
                    "driver_id": str(req.driver.id)
                }
            else:
                return Response(
                    {"error": "Invalid role. Use 'user' or 'driver'"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            return Response(response_data)

        except Request.DoesNotExist:
            return Response(
                {"error": "No active request found"},
                status=status.HTTP_404_NOT_FOUND
            )
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )